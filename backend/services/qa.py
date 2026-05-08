from langchain_community.document_loaders import PyPDFLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_community.vectorstores import FAISS
from langchain_core.documents import Document
from langchain_core.messages import HumanMessage
from langchain_groq import ChatGroq
from services.llm import get_llm
import base64
import os

try:
    import docx
except ImportError:
    docx = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

def extract_docx(path):
    if not docx:
        raise Exception("python-docx is not installed. Please run pip install python-docx")
    doc = docx.Document(path)
    text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
    return [Document(page_content=text)]

def extract_pptx(path):
    if not Presentation:
        raise Exception("python-pptx is not installed. Please run pip install python-pptx")
    prs = Presentation(path)
    text_runs = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text_runs.append(shape.text)
    return [Document(page_content="\n".join(text_runs))]

def extract_image_text(path):
    with open(path, "rb") as f:
        img_b64 = base64.b64encode(f.read()).decode()
    
    # Use vision model
    llm = ChatGroq(
        model_name="llama-3.2-11b-vision-preview",
        api_key=os.getenv("GROQ_API_KEY"),
        temperature=0
    )
    
    # Determine mime type naively
    ext = os.path.splitext(path)[1].lower()
    mime_type = "image/png"
    if ext in [".jpg", ".jpeg"]:
        mime_type = "image/jpeg"
    elif ext == ".webp":
        mime_type = "image/webp"

    msg = HumanMessage(content=[
        {"type": "text", "text": "Extract all text and describe any important diagrams in this image in detail. Be comprehensive."},
        {"type": "image_url", "image_url": {"url": f"data:{mime_type};base64,{img_b64}"}}
    ])
    response = llm.invoke([msg])
    return [Document(page_content=response.content)]

# Pre-load embeddings model at module level for speed
embeddings_model = HuggingFaceEmbeddings(model_name="all-MiniLM-L6-v2")

def build_retriever(path):
    ext = os.path.splitext(path)[1].lower()
    
    if ext == ".pdf":
        loader = PyPDFLoader(path)
        docs = loader.load()
    elif ext in [".doc", ".docx"]:
        docs = extract_docx(path)
    elif ext in [".ppt", ".pptx"]:
        docs = extract_pptx(path)
    elif ext in [".jpg", ".jpeg", ".png", ".webp"]:
        docs = extract_image_text(path)
    else:
        raise ValueError(f"Unsupported file format: {ext}")

    # Optimized chunking strategy for enhanced context retrieval
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=1500, 
        chunk_overlap=150,
        separators=["\n\n", "\n", " ", ""]
    )
    docs = splitter.split_documents(docs)

    db = FAISS.from_documents(docs, embeddings_model)

    return db.as_retriever()

def ask_pdf(retriever, query):
    llm = get_llm()
    docs = retriever.invoke(query)

    context = "\n".join([d.page_content for d in docs])

    prompt = f"""
Answer ONLY from the provided context. If the answer is not in the context, say "I couldn't find information about that in the document."

Format your response using clear paragraphs with double newlines between them. 
Use Markdown for bold text, lists, and headings to make it easy to read.

Context:
{context}

Question: {query}
"""
    res = llm.invoke(prompt)
    return res.content